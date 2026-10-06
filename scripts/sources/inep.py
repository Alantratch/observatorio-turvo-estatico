"""Official INEP file discovery and bounded readers; no invented REST API."""
import csv
import hashlib
import io
import json
import os
import re
import ssl
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen

PORTAL='https://www.gov.br/inep/pt-br/'
PAGES={
 'census':'acesso-a-informacao/dados-abertos/microdados/censo-escolar',
 'ideb':'areas-de-atuacao/pesquisas-estatisticas-e-indicadores/ideb/resultados',
 'saeb':'areas-de-atuacao/avaliacao-e-exames-educacionais/saeb/resultados',
 'literacy':'areas-de-atuacao/avaliacao-e-exames-educacionais/avaliacao-da-alfabetizacao/resultados',
 'synopsis':'acesso-a-informacao/dados-abertos/sinopses-estatisticas/educacao-basica',
 **{k:'acesso-a-informacao/dados-abertos/indicadores-educacionais/'+v for k,v in {
 'afd':'adequacao-da-formacao-docente','ied':'esforco-docente','atu':'media-de-alunos-por-turma',
 'had':'media-de-horas-aula-diaria','dsu':'percentual-de-docentes-com-curso-superior',
 'ird':'regularidade-do-corpo-docente','tdi':'taxas-de-distorcao-idade-serie','flow':'taxas-de-rendimento-escolar'}.items()}}
MUNICIPALITIES={'4127965':'Turvo','4109401':'Guarapuava','4119608':'Pitanga','4113254':'Laranjal'}
NETWORKS={'1':'Federal','2':'Estadual','3':'Municipal','4':'Privada'}
LOCATIONS={'1':'Urbana','2':'Rural'}
SITUATIONS={'1':'Em atividade','2':'Paralisada','3':'Extinta no ano','4':'Extinta em anos anteriores'}

def tls_context():
    context=ssl.create_default_context()
    # INEP's Sep/2026 certificate omits its intermediate. Add the public issuer,
    # never a leaf/root bypass; still require a system-trusted root and hostname.
    context.load_verify_locations(cafile=str(Path(__file__).with_name('certificates')/'rnp-icpedu-gr46-2025.pem'))
    context.verify_flags &= ~getattr(ssl,'VERIFY_X509_PARTIAL_CHAIN',0)
    return context


def request(url,method='GET',timeout=90):
    if urlparse(url).scheme!='https':raise ValueError('Origem INEP exige HTTPS')
    return urlopen(Request(url,method=method,headers={'User-Agent':'ObservatorioTurvo/0.4 (public education data)'}),timeout=timeout,context=tls_context())


class Links(HTMLParser):
    def __init__(self):super().__init__();self.links=[];self.tabs=[];self.href=None;self.label=''
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if a.get('data-url'):self.tabs.append(a['data-url'])
        if tag=='a':self.href=a.get('href');self.label=''
    def handle_data(self,data):
        if self.href:self.label+=data
    def handle_endtag(self,tag):
        if tag=='a' and self.href:self.links.append((self.label.strip(),self.href));self.href=None


def page_links(url):
    with request(url) as r:raw=r.read(2_000_001)
    if len(raw)>2_000_000:raise ValueError('Página de catálogo grande demais')
    p=Links();p.feed(raw.decode('utf-8'));return p


def year_of(text):
    years=re.findall(r'(?<!\d)(20\d{2})(?!\d)',text)
    return max(map(int,years)) if years else None


def discover(year=None,history_start=2015):
    """Only catalog HTML, never national downloads. Every URL comes from INEP."""
    manifest=os.environ.get('INEP_MANIFEST_URL')
    if manifest:
        with request(manifest) as r:return json.load(r)
    def one(item):
        key,path=item;official=PORTAL+path;p=page_links(official);links=list(p.links)
        tabs=[u for u in p.tabs if u.startswith(PORTAL) and year_of(u)]
        if tabs:
            desired=max(year_of(u) for u in tabs) if year is None or key in ('ideb','saeb','literacy') else year
            selected=[u for u in tabs if year_of(u)==desired]
            if not selected:raise ValueError(f'{key}: ano {desired} não localizado no catálogo')
            for u in selected:links+=page_links(u).links
        files=[]
        for label,url in links:
            if urlparse(url).hostname!='download.inep.gov.br' or not url.lower().endswith(('.zip','.xlsx')):continue
            ref=year_of(url)
            if not ref:continue
            if key in ('census','synopsis'):
                if ref<history_start or (year and ref>year):continue
            elif key in ('ideb','saeb','literacy'):
                pass # Assessment cycles are independently discovered, not forced to census year.
            elif 'municip' not in (url+' '+label).lower():continue
            files.append({'key':key,'reference':str(ref),'url':url,'officialUrl':official,'label':label})
        if key not in ('census','synopsis') and files:
            latest=max(int(f['reference']) for f in files);files=[f for f in files if int(f['reference'])==latest]
        if key=='ideb':files=[f for f in files if 'municipios' in f['url'] or 'regioes_ufs_ideb' in f['url']]
        if key=='literacy':files=[f for f in files if 'municipios' in f['url']]
        if key=='synopsis' and files:files=[max(files,key=lambda f:int(f['reference']))]
        if not files:raise ValueError(f'{key}: nenhum arquivo oficial localizado')
        return key,files
    return dict(ThreadPoolExecutor(6).map(one,PAGES.items()))


def file_hash(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        while chunk:=f.read(262144):h.update(chunk)
    return h.hexdigest()


def download(url,path,force=False,attempts=3):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);stamp=path.with_suffix(path.suffix+'.manifest.json')
    if path.exists() and stamp.exists():
        try:
            meta=json.loads(stamp.read_text());valid=meta['url']==url and meta['bytes']==path.stat().st_size and meta['sha256']==file_hash(path)
            if valid and not force:return meta
            if valid:
                with request(url,'HEAD',30) as r:remote={k:r.headers.get(k) for k in ('Content-Length','ETag','Last-Modified')}
                if (remote.get('ETag') or remote.get('Last-Modified')) and remote==meta.get('remote'):return meta
        except Exception:pass
    partial=path.with_suffix(path.suffix+'.part')
    for attempt in range(attempts):
        try:
            size=0;h=hashlib.sha256()
            with request(url) as r,partial.open('wb') as f:
                remote={k:r.headers.get(k) for k in ('Content-Length','ETag','Last-Modified')};expected=int(remote['Content-Length'] or 0)
                while chunk:=r.read(262144):f.write(chunk);h.update(chunk);size+=len(chunk)
            if size<32 or (expected and size!=expected):raise ValueError('Download incompleto')
            with zipfile.ZipFile(partial) as z:
                if not z.namelist():raise ValueError('Arquivo ZIP/XLSX vazio')
            meta={'url':url,'bytes':size,'sha256':h.hexdigest(),'remote':remote}
            os.replace(partial,path);stamp.write_text(json.dumps(meta)+'\n');return meta
        except Exception:
            partial.unlink(missing_ok=True)
            if attempt==attempts-1:raise
            time.sleep(2**attempt)


def workbooks(path):
    """Read-only openpyxl: handles actual merged/multirow INEP layouts."""
    from openpyxl import load_workbook
    path=Path(path)
    if path.suffix.lower() in ('.xlsx','.xlsm'):
        yield path.name,load_workbook(path,read_only=True,data_only=True)
    else:
        with zipfile.ZipFile(path) as z:
            names=[n for n in z.namelist() if n.lower().endswith(('.xlsx','.xlsm')) and not '/~$' in n]
            if not names:raise ValueError('Nenhuma planilha XLSX oficial encontrada')
            for name in names:
                yield name,load_workbook(io.BytesIO(z.read(name)),read_only=True,data_only=True)


def census_rows(path,year,codes=None):
    codes=set(codes or MUNICIPALITIES)
    with zipfile.ZipFile(path) as z:
        legacy=[n for n in z.namelist() if n.lower().endswith('.csv') and f'microdados_ed_basica_{year}' in n.lower()]
        split={kind:[n for n in z.namelist() if n.lower().endswith('.csv') and f'tabela_{kind}_{year}' in n.lower()] for kind in ('escola','matricula','docente','turma')}
        files=legacy if len(legacy)==1 else [values[0] for values in split.values() if len(values)==1]
        if len(files) not in (1,4):raise ValueError('Layout escolar ausente/ambíguo; esperado CSV legado ou quatro tabelas escolares')
        combined={};parts={}
        for member in files:
            seen=set()
            with io.TextIOWrapper(z.open(member),encoding='cp1252',errors='strict') as stream:
                reader=csv.DictReader(stream,delimiter=';');required={'NU_ANO_CENSO','CO_MUNICIPIO','SG_UF','CO_ENTIDADE','NO_ENTIDADE','TP_DEPENDENCIA','TP_LOCALIZACAO'}
                expected='QT_MAT_BAS' if len(files)==1 or 'tabela_matricula_' in member.lower() else 'QT_DOC_BAS' if 'tabela_docente_' in member.lower() else 'QT_TUR_BAS' if 'tabela_turma_' in member.lower() else 'TP_SITUACAO_FUNCIONAMENTO'
                required.add(expected)
                if not reader.fieldnames or len(reader.fieldnames)!=len(set(reader.fieldnames)) or not required.issubset(reader.fieldnames):raise ValueError('Layout Censo Escolar incompatível')
                for row in reader:
                    if None in row or any(v is None for v in row.values()):raise ValueError('Registro escolar truncado')
                    if row['CO_MUNICIPIO'] not in codes:continue
                    if row['SG_UF']!='PR' or row['NU_ANO_CENSO']!=str(year):raise ValueError('UF/ano divergente')
                    identifier=row['CO_ENTIDADE']
                    if identifier in seen:raise ValueError('Escola duplicada dentro da tabela oficial')
                    seen.add(identifier);base=combined.setdefault(identifier,{})
                    for key,value in row.items():
                        if key in base and base[key]!=value:raise ValueError(f'Junção escolar divergente: {identifier} {key}')
                        base[key]=value
            parts[expected]=seen
        for identifier,row in combined.items():
            if row.get('TP_SITUACAO_FUNCIONAMENTO')=='1' and row.get('QT_MAT_BAS') not in (None,'') and any(identifier not in seen for seen in parts.values()):raise ValueError('Escola ativa sem registro em todas as tabelas')
            yield row
