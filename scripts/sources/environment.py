"""Public environmental sources; bounded downloads and small municipal extracts."""

import csv
import hashlib
import html
import io
import json
import math
import re
import tempfile
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from scripts.common import request_json, atomic_write, read

MUNICIPALITIES = {
    "4127965": "Turvo",
    "4109401": "Guarapuava",
    "4119608": "Pitanga",
    "4113254": "Laranjal",
}
MAP_PORTAL = "https://brasil.mapbiomas.org/downloads/estatisticas/"
FIRE_BASE = "https://dataserver-coids.inpe.br/queimadas/queimadas/focos/csv/"
FIRE_PORTAL = "https://data.inpe.br/queimadas/pages/secao_downloads/dados-abertos/"
ANA = "https://portal1.snirh.gov.br/server/rest/services/Esta%C3%A7%C3%B5es_Hidrometeorol%C3%B3gicas_SNIRH/FeatureServer/0"
UC = "https://pamgia.ibama.gov.br/server/rest/services/BasesSincronizadas/lim_unidades_conserva%C3%A7%C3%A3o_mma_a/FeatureServer/0"
BIOME = "https://pamgia.ibama.gov.br/server/rest/services/BasesSincronizadas/lim_biomas_ibge_250_a/FeatureServer/0"
SATELLITE = "AQUA_M-T"
NATIVE_CODES = {3, 4, 5, 6, 7, 49, 11, 12, 77, 84, 50}


def text(value):
    return (
        "".join(
            c
            for c in unicodedata.normalize("NFKD", str(value))
            if not unicodedata.combining(c)
        )
        .strip()
        .upper()
    )


def source(identifier, agency, dataset, url, reference, now, unit, **extra):
    return dict(
        id=identifier,
        agency=agency,
        dataset=dataset,
        url=url,
        reference=reference,
        collectedAt=now,
        unit=unit,
        municipalityCodes=extra.pop("municipalityCodes", list(MUNICIPALITIES)),
        **extra,
    )


def request(url, method="GET"):
    if not url.startswith("https://"):
        raise ValueError("Fonte deve usar HTTPS")
    return urlopen(
        Request(
            url,
            method=method,
            headers={"User-Agent": "ObservatorioTurvo/1.0 (public environmental data)"},
        ),
        timeout=30,
    )


def download(url, path, limit=160_000_000):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    count = 0
    try:
        with request(url) as response, path.with_suffix(path.suffix + ".part").open(
            "wb"
        ) as out:
            expected = response.headers.get("Content-Length")
            if expected and int(expected) > limit:
                raise ValueError("Arquivo excede limite de download")
            while chunk := response.read(1024 * 1024):
                count += len(chunk)
                if count > limit:
                    raise ValueError("Arquivo excede limite de download")
                digest.update(chunk)
                out.write(chunk)
            if expected and count != int(expected):
                raise ValueError("Download incompleto")
        path.with_suffix(path.suffix + ".part").replace(path)
    finally:
        path.with_suffix(path.suffix + ".part").unlink(missing_ok=True)
    return digest.hexdigest()


def valid_number(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("Número ausente")
    result = float(value)
    if not math.isfinite(result) or result < 0:
        raise ValueError("Área inválida")
    return result


def parse_land(rows, legend, collection):
    selected = []
    seen = set()
    for row in rows:
        code = str(row["geocode"])
        if code not in MUNICIPALITIES:
            continue
        if text(row["state"]) != "PARANA" or text(row["municipality"]) != text(
            MUNICIPALITIES[code]
        ):
            raise ValueError("Identificação municipal MapBiomas divergente")
        identifier = int(row["class"])
        label = legend.get(str(identifier))
        if not label:
            raise ValueError("Classe sem legenda oficial")
        key = (code, str(row["biome"]), identifier)
        if key in seen:
            raise ValueError("Linha municipal/classe/bioma duplicada")
        seen.add(key)
        match = re.match(r"([\d.]+)\s", label["name"])
        if not match:
            raise ValueError("Hierarquia da legenda ausente")
        prefix = match[1].rstrip(".")
        years = {
            key[1:]: valid_number(value)
            for key, value in row.items()
            if re.fullmatch(r"y\d{4}", str(key))
        }
        if not years or sorted(years) != [
            str(y) for y in range(1985, int(max(years)) + 1)
        ]:
            raise ValueError("Série de cobertura incompleta")
        selected.append(
            {
                "municipalityCode": code,
                "biome": row["biome"],
                "code": identifier,
                "name": label["name"],
                "level": len(prefix.split(".")),
                "hierarchy": prefix,
                "color": label["color"],
                "collection": collection,
                "unit": "ha",
                "years": years,
            }
        )
    if set(r["municipalityCode"] for r in selected) != set(MUNICIPALITIES):
        raise ValueError("Municípios incompletos no MapBiomas")
    return selected


def land(path, collection):
    from openpyxl import load_workbook

    # MapBiomas publishes a ZIP containing one XLSX, not a raw XLSX.
    with tempfile.TemporaryDirectory(prefix="turvo-mapbiomas-") as folder:
        with zipfile.ZipFile(path) as archive:
            if "[Content_Types].xml" in archive.namelist():
                workbook_path = path
            else:
                items = [
                    i
                    for i in archive.infolist()
                    if i.filename.lower().endswith(".xlsx")
                ]
                if len(items) != 1 or items[0].file_size > 160_000_000:
                    raise ValueError("ZIP MapBiomas inválido")
                workbook_path = Path(folder) / "data.xlsx"
                workbook_path.write_bytes(archive.read(items[0]))
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            required = {"COVERAGE_" + collection, "LEGEND_CODE"}
            if not required.issubset(workbook.sheetnames):
                raise ValueError("Coleção/layout não confirmado")
            legend = {
                str(int(r[3])): {"name": r[1], "color": r[4]}
                for r in workbook["LEGEND_CODE"].iter_rows(values_only=True)
                if isinstance(r[3], (int, float)) and r[1]
            }
            iterator = workbook["COVERAGE_" + collection].iter_rows(values_only=True)
            header = next(iterator)
            if tuple(header[:9]) != (
                "ID",
                "country",
                "biome",
                "region",
                "state",
                "geocode",
                "municipality",
                "municipality-state",
                "class",
            ):
                raise ValueError("Layout MapBiomas mudou")
            return parse_land(
                (dict(zip(header, row)) for row in iterator), legend, collection
            )
        finally:
            workbook.close()


def discover_land():
    with request(MAP_PORTAL) as response:
        page = response.read(2_000_000).decode("utf-8")
    rows = re.findall(r"<tr\b[^>]*>(.*?)</tr>", page, flags=re.S)
    matches = []
    for row in rows:
        if "Biomas, Estados e Municípios | Cobertura 30m" not in row:
            continue
        version = re.search(r"Coleção\s+(\d+)", row)
        links = re.findall(r'href="([^"]+)"', row)
        downloads = [
            html.unescape(link)
            for link in links
            if "drive.google.com/uc?" in link or link.endswith(".xlsx")
        ]
        if version and downloads:
            matches.append((int(version[1]), downloads[0]))
    if not matches:
        raise ValueError("Tabulação municipal não encontrada no catálogo MapBiomas")
    collection, url = max(matches)
    if "drive.google.com" in url:
        identifier = re.search(r"[?&]id=([\w-]+)", url)
        if not identifier:
            raise ValueError("Link público MapBiomas inválido")
        download_url = "https://drive.usercontent.google.com/download?" + urlencode(
            {"id": identifier[1], "export": "download", "confirm": "t"}
        )
    else:
        download_url = url
    return str(collection), url, download_url


def parse_fire(rows, reference, annual=False):
    results = []
    seen = {}
    latest = None
    names = {text(name): code for code, name in MUNICIPALITIES.items()}
    required = (
        {"lat", "lon", "estado", "municipio", "foco_id", "data_pas"}
        if annual
        else {
            "id",
            "lat",
            "lon",
            "municipio_id",
            "estado",
            "municipio",
            "data_hora_gmt",
            "satelite",
        }
    )
    for row in rows:
        if not required.issubset(row):
            raise ValueError("Layout INPE mudou")
        date = row["data_pas" if annual else "data_hora_gmt"].strip()
        instant = datetime.fromisoformat(date).replace(tzinfo=timezone.utc)
        if not date.startswith(reference[:4]) or (
            not annual and not date.startswith(reference[:4] + "-" + reference[4:])
        ):
            raise ValueError("Data fora do arquivo de referência")
        if instant > datetime.now(timezone.utc):
            raise ValueError("Data INPE futura")
        latest = max(latest or date, date)
        if not annual and row["satelite"].strip() != SATELLITE:
            continue
        if text(row["estado"]) != "PARANA":
            continue
        if annual:
            code = names.get(text(row["municipio"]))
        else:
            code = str(row["municipio_id"]).strip()
            if code in MUNICIPALITIES and text(row["municipio"]) != text(
                MUNICIPALITIES[code]
            ):
                raise ValueError("Nome/código INPE divergentes")
        if code not in MUNICIPALITIES:
            continue
        lat, lon = float(row["lat"]), float(row["lon"])
        if not math.isfinite(lat + lon) or not (-27 < lat < -22 and -55 < lon < -48):
            raise ValueError("Coordenadas INPE inválidas para PR")
        identifier = row["foco_id" if annual else "id"].strip()
        if not identifier:
            raise ValueError("Identificador INPE ausente")
        # Same official ID repeated is one observation; distinct detections at one location survive.
        observation = {
            "id": identifier,
            "municipalityCode": code,
            "date": instant.isoformat(),
            "satellite": SATELLITE,
            "latitude": lat,
            "longitude": lon,
            "biome": row.get("bioma") or None,
        }
        if identifier in seen:
            if seen[identifier] != observation:
                raise ValueError("ID INPE conflitante")
            continue
        seen[identifier] = observation
        results.append(observation)
    if latest is None:
        raise ValueError("Arquivo INPE vazio; ausência não pode virar zero")
    return {
        "observations": results,
        "availableThrough": latest[:10],
        "reference": reference,
        "coverage": "annual-complete" if annual else "monthly-export",
        "satellite": SATELLITE,
    }


def fire(path, reference, annual=False):
    if annual:
        with zipfile.ZipFile(path) as archive:
            items = [i for i in archive.infolist() if i.filename.endswith(".csv")]
            if len(items) != 1 or items[0].file_size > 160_000_000:
                raise ValueError("Arquivo INPE inválido")
            with archive.open(items[0]) as handle:
                return parse_fire(
                    csv.DictReader(io.TextIOWrapper(handle, encoding="utf-8-sig")),
                    reference,
                    True,
                )
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        return parse_fire(csv.DictReader(handle), reference)


def geo_query(base, bounds, fields="*", where="1=1"):
    params = {
        "f": "geojson",
        "where": where,
        "geometry": ",".join(map(str, bounds)),
        "geometryType": "esriGeometryEnvelope",
        "inSR": 4326,
        "spatialRel": "esriSpatialRelIntersects",
        "outFields": fields,
        "outSR": 4326,
        "returnGeometry": "true",
        "resultRecordCount": 2000,
    }
    url = base + "/query?" + urlencode(params)
    result = request_json(url, timeout=60, attempts=2)
    if (
        result.get("error")
        or result.get("exceededTransferLimit")
        or len(result.get("features", [])) >= 2000
        or not isinstance(result.get("features"), list)
    ):
        raise ValueError("Resposta geográfica incompleta")
    return result, url


def stations(features, boundary):
    from shapely.geometry import Point

    centroid = boundary.centroid
    results = []
    seen = set()
    for feature in features:
        p = feature["properties"]
        code = str(p["Codigo"]).zfill(8)
        if code in seen:
            raise ValueError("Estação ANA duplicada")
        seen.add(code)
        if text(p["UF"]) not in {"PARANA", "PR"}:
            continue
        lon, lat = feature["geometry"]["coordinates"][:2]
        if (
            not math.isfinite(lat + lon)
            or abs(lat - float(p["Latitude"])) > 0.0001
            or abs(lon - float(p["Longitude"])) > 0.0001
        ):
            raise ValueError("Coordenadas ANA divergentes")
        if p["TipoEstacao"] not in {"Fluviométrica", "Pluviométrica"}:
            raise ValueError("Tipo ANA não confirmado")
        distance = (
            6371
            * 2
            * math.asin(
                min(
                    1,
                    math.sqrt(
                        math.sin(math.radians(lat - centroid.y) / 2) ** 2
                        + math.cos(math.radians(lat))
                        * math.cos(math.radians(centroid.y))
                        * math.sin(math.radians(lon - centroid.x) / 2) ** 2
                    ),
                )
            )
        )
        inside = boundary.covers(Point(lon, lat))
        if not inside and distance > 50:
            continue
        results.append(
            {
                "code": code,
                "name": p["Nome"],
                "type": p["TipoEstacao"],
                "latitude": lat,
                "longitude": lon,
                "location": "inside" if inside else "regional",
                "sourceMunicipality": p["Municipio"],
                "sourceMunicipalityCode": p["MunicipioCodigo"],
                "distanceToCentroidKm": distance,
                "operating": p["Operando"],
                "telemetric": p["EstacaoTelemetrica"],
                "agency": p["ResponsavelSigla"],
                "basin": p["Bacia"],
                "subBasin": p["SubBacia"],
                "river": p["Rio"] or None,
                "registryUpdatedAt": p["DataAlteracao"],
            }
        )
    inside = sorted(
        [r for r in results if r["location"] == "inside"], key=lambda r: r["code"]
    )
    nearby = sorted(
        [r for r in results if r["location"] == "regional"],
        key=lambda r: (r["distanceToCentroidKm"], r["code"]),
    )[:10]
    return inside + nearby


def protected(features, boundary, area_km2):
    from shapely.geometry import shape, mapping
    from shapely.ops import transform, unary_union
    from pyproj import Transformer

    project = Transformer.from_crs(4326, 31982, always_xy=True).transform
    polygons = []
    units = []
    seen = set()
    for feature in features:
        p = feature["properties"]
        code = str(p["cd_cnuc"])
        if code in seen:
            raise ValueError("UC duplicada")
        seen.add(code)
        geometry = shape(feature["geometry"])
        if not geometry.is_valid:
            raise ValueError("Polígono UC inválido")
        clipped = geometry.intersection(boundary)
        area = transform(project, clipped).area / 10000
        if area <= 0:
            continue
        polygons.append(clipped)
        units.append(
            {
                "code": code,
                "name": p["nome_uc"],
                "sphere": p["esfera"],
                "category": p["categoria"],
                "group": p["grupo"],
                "agency": p["org_gestor"],
                "listedMunicipalities": p["municipio"],
                "registryMentionsTurvo": bool(
                    re.search(r"(?<![A-ZÀ-Ú])TURVO\s*\(PR\)", text(p["municipio"]))
                ),
                "intersectionHa": area,
                "syncedAt": p["dt_sync"],
                "geometry": mapping(clipped.simplify(0.00002, preserve_topology=True)),
            }
        )
    union = transform(project, unary_union(polygons)).area / 10000 if polygons else 0
    if union > area_km2 * 100:
        raise ValueError("Interseção de UCs excede área municipal")
    return {
        "status": "derived",
        "units": sorted(units, key=lambda r: r["code"]),
        "unionAreaHa": union,
        "sharePercent": union / (area_km2 * 100) * 100,
        "declaredMunicipalityCount": sum(u["registryMentionsTurvo"] for u in units),
        "projection": "EPSG:31982",
        "methodology": "União das interseções com malha IBGE 2022; área projetada SIRGAS 2000 / UTM 22S. Estimativa cartográfica, sem dupla contagem; não equivale a demarcação legal, APP, RL ou CAR.",
    }


def fire_items(now):
    """Discover annual archives; use monthly exports when the annual file is pending."""

    def years(folder, pattern):
        with request(FIRE_BASE + folder) as response:
            page = response.read(2_000_000).decode("utf-8")
        return {int(y) for y in re.findall(pattern, page)}

    state = years("anual/EstadosBr_sat_ref/PR/", r"focos_br_pr_ref_(\d{4})\.zip")
    national = years("anual/Brasil_sat_ref/", r"focos_br_ref_(\d{4})\.zip")
    items = []
    for year in range(2020, now.year + 1):
        if year < now.year and year in state:
            items.append(
                (
                    str(year),
                    FIRE_BASE
                    + f"anual/EstadosBr_sat_ref/PR/focos_br_pr_ref_{year}.zip",
                    True,
                )
            )
        elif year < now.year and year in national:
            items.append(
                (
                    str(year),
                    FIRE_BASE + f"anual/Brasil_sat_ref/focos_br_ref_{year}.zip",
                    True,
                )
            )
        else:
            last = now.month if year == now.year else 12
            items.extend(
                (
                    f"{year}{month:02d}",
                    FIRE_BASE + f"mensal/Brasil/focos_mensal_br_{year}{month:02d}.csv",
                    False,
                )
                for month in range(1, last + 1)
            )
    return items
