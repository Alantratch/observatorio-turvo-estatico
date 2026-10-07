"""Public MDS municipal aggregates. Never acquire identified CadÚnico/RH data.

The Solr service is used by RI/VIS DATA; six-digit codes are mapped explicitly.
Only the Censo catalogue (resource discovery) uses HTML, not indicator scraping.
"""

import csv
import hashlib
import html
import io
import json
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlencode, urlparse, unquote
from scripts.common import request_json
from scripts.sources.datasus import download

MUNICIPALITIES = {
    "4127965": "Turvo",
    "4119608": "Pitanga",
    "4113254": "Laranjal",
    "4109401": "Guarapuava",
}
SERVICE = "https://aplicacoes.mds.gov.br/sagi/servicos/misocial"
CENSO = "https://aplicacoes.mds.gov.br/sagi/snas/vigilancia/index2"
REPORT = "https://aplicacoes.cidadania.gov.br/ri/ri/relatorios/cidadania/"
CAD = {
    "families": "cadun_qtd_familias_cadastradas_i",
    "people": "cadun_qtd_pessoas_cadastradas_i",
    "updated": "cadun_qtd_familias_atualizadas_i",
    "updatedPercent": "cadun_taxa_atualizacao_cadastral_d",
    "povertyFamilies": "cadun_qtd_familias_cadastradas_pobreza_pbf_i",
    "lowIncomeFamilies": "cadun_qtd_familias_cadastradas_baixa_renda_i",
    "aboveHalfFamilies": "cadun_qtd_familias_cadastradas_rfpc_acima_meio_sm_i",
    "povertyPeople": "cadun_qtd_pessoas_cadastradas_pobreza_pbf_i",
    "lowIncomePeople": "cadun_qtd_pessoas_cadastradas_baixa_renda_i",
    "aboveHalfPeople": "cadun_qtd_pessoas_cadastradas_rfpc_acima_meio_sm_i",
}
BF = {
    "families": "qtd_familias_beneficiarias_bolsa_familia",
    "people": "qtd_pessoas_beneficiarias_bolsa_familia_i",
    "transferredValue": "valor_repassado_bolsa_familia",
    "averageBenefit": "pbf_vlr_medio_benef_f",
}
BPC = {
    "total": "bpc_ben_i",
    "elderly": "bpc_idoso_ben_i",
    "disabled": "bpc_pcd_ben_i",
    "elderlyValue": "bpc_idoso_val_s",
    "disabledValue": "bpc_pcd_val_s",
    "transferredValue": "bpc_val_s",
}
NETWORK = {
    "cras": "cadsuas_qtd_cras_i",
    "creas": "cadsuas_qtd_creas_i",
    "centroPop": "cadsuas_qtd_centro_pop_i",
    "acolhimento": "cadsuas_qtd_unidade_acolhimento_i",
    "convivencia": "cadsuas_qtd_centro_convivencia_i",
    "centroDia": "cadsuas_qtd_centro_dia_similares_i",
}
TYPES = {
    "1": "cras",
    "2": "creas",
    "3": "centroPop",
    "4": "acolhimento",
    "5": "convivencia",
    "6": "centroDia",
    "8": "cadunicoPost",
    "13": "council",
}
LABELS = {
    "cras": "CRAS",
    "creas": "CREAS",
    "centroPop": "Centro POP",
    "acolhimento": "Unidade de Acolhimento",
    "convivencia": "Centro de Convivência",
    "centroDia": "Centro Dia",
    "cadunicoPost": "Posto do Cadastro Único",
    "council": "Conselho Municipal de Assistência Social",
}
IDENTITY = {
    "codigo_ibge",
    "sigla_uf",
    "municipio",
    "anomes_s",
    "tipo_s",
    "cadsuas_data_extracao_s",
}
FIELDS = (
    IDENTITY
    | set(CAD.values())
    | set(BF.values())
    | set(BPC.values())
    | set(NETWORK.values())
)


def numeric(value, money=False):
    if value is None or str(value).strip() in {"", "-", "..", "...", "X"}:
        return None
    # Source uses decimal dot; accepting BR formatting here would conceal schema changes.
    try:
        n = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError("Número MDS inválido") from None
    if not n.is_finite() or n < 0:
        raise ValueError("Valor MDS negativo/não finito")
    if money:
        return float(n.quantize(Decimal("0.01")))
    if n != n.to_integral():
        raise ValueError("Contagem não inteira")
    return int(n)


def percent(value):
    if value is None:
        return None
    n = float(Decimal(str(value)))
    if not 0 <= n <= 100:
        raise ValueError("Percentual fora de 0–100")
    return n


def valid_month(value):
    return bool(re.fullmatch(r"20\d{2}(0[1-9]|1[0-2])", str(value)))


def monthly_url(start, end):
    return (
        SERVICE
        + "?"
        + urlencode(
            {
                "q": "*:*",
                "fq": [
                    "tipo_s:mes_mu",
                    "codigo_ibge:(" + " OR ".join(c[:6] for c in MUNICIPALITIES) + ")",
                    f"anomes_s:[{start} TO {end}]",
                ],
                "wt": "json",
                "rows": 400,
                "sort": "anomes_s asc",
                "fl": ",".join(sorted(FIELDS)),
            },
            doseq=True,
        )
    )


def network_url(extracted_at):
    return (
        SERVICE
        + "?"
        + urlencode(
            {
                "q": "*:*",
                "fq": [
                    "tipo_s:mes_mu",
                    "codigo_ibge:412796",
                    f'cadsuas_data_extracao_s:"{extracted_at}"',
                ],
                "wt": "json",
                "rows": 1,
                "sort": "anomes_s desc",
                "fl": ",".join(sorted(IDENTITY | set(NETWORK.values()))),
            },
            doseq=True,
        )
    )


def monthly(start, end):
    if not valid_month(start) or not valid_month(end) or start > end:
        raise ValueError("Janela mensal inválida")
    url = monthly_url(start, end)
    response = request_json(url)["response"]
    docs = response["docs"]
    if response["numFound"] != len(docs) or response.get("start", 0) != 0:
        raise ValueError("Resposta MDS truncada")
    by_code = {c: [] for c in MUNICIPALITIES}
    seen = set()
    for original in docs:
        r = {k: v for k, v in original.items() if k in FIELDS}
        code = next(
            (c for c in MUNICIPALITIES if c[:6] == str(r.get("codigo_ibge"))), None
        )
        p = r.get("anomes_s")
        key = (code, p)
        if (
            not code
            or r.get("sigla_uf") != "PR"
            or str(r.get("municipio", "")).casefold() != MUNICIPALITIES[code].casefold()
            or r.get("tipo_s") != "mes_mu"
        ):
            raise ValueError("Território MDS divergente")
        if not valid_month(p) or not start <= p <= end or key in seen:
            raise ValueError("Competência inválida/duplicada")
        seen.add(key)
        by_code[code].append(r)
    if any(not rows for rows in by_code.values()):
        raise ValueError("Município não retornado pelo MDS")
    return by_code, url


def latest(rows, fields, earliest="202303"):
    complete = [
        r
        for r in rows
        if r["anomes_s"] >= earliest
        and all(r.get(k) is not None for k in fields.values())
    ]
    return max(complete, key=lambda r: r["anomes_s"]) if complete else None


def resources(text):
    """Only accept links actually advertised by the official catalogue."""
    links = [
        html.unescape(u) for u in re.findall(r'href=["\']([^"\']+)["\']', text, re.I)
    ]
    censos = {}
    rma = []
    for u in links:
        p = urlparse(u)
        if p.scheme != "https" or p.netloc != "aplicacoes.mds.gov.br":
            continue
        m = re.search(r"/censosuas/(20\d{2})/(\d+)_.*\.zip$", p.path, re.I)
        if m and m[2] in TYPES:
            censos.setdefault(m[1], {})[TYPES[m[2]]] = u
        if re.search(r"RMA_CRAS_.*_(20\d{2})_.*\.xlsx$", unquote(p.path), re.I):
            rma.append((re.search(r"_(20\d{2})_", unquote(p.path))[1], u))
    if not censos:
        raise ValueError("Catálogo Censo SUAS não apresenta arquivos compatíveis")
    year = max(censos)
    return year, censos[year], max(rma) if rma else None


def acquire(url, cache, force=False):
    """ETag/Last-Modified/length revalidation; unconditional refresh if absent.
    A sidecar records HTTP validators, never national data in public/data.
    """
    from urllib.request import Request, urlopen

    cache = Path(cache)
    cache.mkdir(parents=True, exist_ok=True)
    path = cache / (
        hashlib.sha256(url.encode()).hexdigest()[:20] + Path(urlparse(url).path).suffix
    )
    side = path.with_suffix(path.suffix + ".json")
    with urlopen(Request(url, method="HEAD"), timeout=40) as response:
        tags = {
            k: response.headers.get(k)
            for k in ("ETag", "Last-Modified", "Content-Length")
        }
    old = json.loads(side.read_text()) if side.exists() else None
    if (
        not force
        and path.exists()
        and old
        and (tags["ETag"] or tags["Last-Modified"])
        and old["validators"] == tags
        and old["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    ):
        return path, old
    detail = download(url, path, 120_000_000)
    detail["validators"] = tags
    detail["url"] = url
    side.write_text(json.dumps(detail))
    return path, detail


def census_rows(path):
    # RH members are intentionally never opened, including their dictionaries.
    with zipfile.ZipFile(path) as archive:
        csvs = [
            m
            for m in archive.infolist()
            if "Dados_Gerais" in m.filename and m.filename.endswith(".csv")
        ]
        xlsx = [
            m
            for m in archive.infolist()
            if "Dados_Gerais" in m.filename and m.filename.endswith(".xlsx")
        ]
        chosen = csvs or xlsx
        if len(chosen) != 1 or chosen[0].file_size > 120_000_000:
            raise ValueError("Censo SUAS: arquivo institucional ausente/ambíguo")
        raw = archive.read(chosen[0])
        if csvs:
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin1")
            yield from csv.DictReader(io.StringIO(text), delimiter=";")
        else:
            import openpyxl

            # Some official XLSX contain XML 1.0-forbidden control bytes. Remove
            # only those bytes, retaining every cell value and original archive hash.
            output = io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(raw)) as original, zipfile.ZipFile(
                output, "w"
            ) as fixed:
                for member in original.infolist():
                    content = original.read(member)
                    if member.filename.endswith(".xml"):
                        content = re.sub(rb"[\x00-\x08\x0b\x0c\x0e-\x1f]", b"", content)
                    fixed.writestr(member, content)
            output.seek(0)
            workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
            if any(
                next(s.iter_rows(values_only=True), None) is not None
                for s in workbook.worksheets[1:]
            ):
                raise ValueError("Planilhas institucionais ambíguas")
            it = workbook.worksheets[0].iter_rows(values_only=True)
            headers = next(it)
            for row in it:
                yield dict(zip(headers, row))
            workbook.close()


def institutional(rows, kind, reference, source_id):
    if reference != "2025":
        raise ValueError(
            "Dicionário institucional Censo SUAS ainda não validado para esta edição"
        )
    result = {c: [] for c in MUNICIPALITIES}
    seen = {}
    for r in rows:
        code = str(r.get("IBGE7", "")).strip()
        if code not in MUNICIPALITIES:
            # Missing 7-digit code does not grant an implicit correspondence.
            if str(r.get("IBGE", "")).strip() in {c[:6] for c in MUNICIPALITIES}:
                raise ValueError("Censo sem IBGE7 confirmado")
            continue
        if str(r.get("IBGE", "")).strip() != code[:6] or (
            str(r.get("UF", "")).strip() not in {"PR", "Paraná"}
            and not (kind == "centroDia" and str(r.get("q0_9", "")).strip() == "PR")
        ):
            raise ValueError("Censo: UF/código divergentes")
        identifier = str(r.get("NU_IDENTIFICADOR", "")).strip()
        if not identifier or not re.fullmatch(r"\d+", identifier):
            raise ValueError("Unidade sem identificador institucional")
        services = []
        # Dictionary 2025: CRAS q12_13=acompanhamento de famílias PAIF;
        # CREAS q12_3=acompanhamento individual/familiar PAEFI.
        if reference == "2025" and kind in {"cras", "creas"}:
            key = "q12_13" if kind == "cras" else "q12_3"
            if str(r.get(key, "")).strip() == "Sim":
                services = ["PAIF" if kind == "cras" else "PAEFI"]
        item = {
            "id": identifier,
            "type": kind,
            "institutionName": str(r.get("q0_1", "")).strip(),
            "institutionalAddress": ", ".join(
                str(r.get(k, "")).strip()
                for k in ["q0_2", "q0_3", "q0_4", "q0_6"]
                if str(r.get(k, "")).strip()
            ),
            "municipalityCode": code,
            "state": "PR",
            "reference": reference,
            "sourceId": source_id,
            "situation": f"Respondente ao Censo SUAS {reference}; não comprova situação operacional atual.",
            "services": services,
            "coordinates": None,
            "coordinateNote": "Coordenadas não utilizadas: exportação sem precisão/validade geográfica confirmada.",
        }
        key = (code, identifier)
        if key in seen:
            if seen[key] != item:
                raise ValueError("Unidade duplicada com atributos divergentes")
            continue
        if not item["institutionName"]:
            raise ValueError("Identificação institucional vazia")
        seen[key] = item
        result[code].append(item)
    for a in result.values():
        a.sort(key=lambda u: u["id"])
    return result


def paif(path, year):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Base tratada" not in workbook.sheetnames:
        raise ValueError("RMA precisa da Base tratada oficial")
    it = workbook["Base tratada"].iter_rows(values_only=True)
    headers = next(it)
    required = {
        "IBGE",
        "IBGE7",
        "UF_A",
        "ano",
        "mes",
        "NU_IDENTIFICADOR",
        "a1",
        "a2",
        "c1",
    }
    if not required.issubset(headers):
        raise ValueError("Schema RMA divergente")
    grouped = {}
    seen = set()
    for row in it:
        r = dict(zip(headers, row))
        code = str(r.get("IBGE7", ""))
        if code != "4127965":
            continue
        if (
            str(r["IBGE"]) != code[:6]
            or r["UF_A"] != "PR"
            or int(r["ano"]) != int(year)
        ):
            raise ValueError("Território/ano RMA divergente")
        month = int(r["mes"])
        p = f"{year}{month:02d}"
        key = (str(r["NU_IDENTIFICADOR"]), p)
        if not valid_month(p) or key in seen:
            raise ValueError("Unidade/mês RMA duplicado ou inválido")
        seen.add(key)
        grouped.setdefault(p, []).append({k: numeric(r[k]) for k in ("a1", "a2", "c1")})
    workbook.close()
    # Missing cells in any reporting unit invalidate municipal total; no imputation.
    return [
        {
            "reference": f"{year}{m:02d}",
            "reportingUnits": len(grouped.get(f"{year}{m:02d}", [])),
            **{
                k: (
                    None
                    if not grouped.get(f"{year}{m:02d}")
                    or any(r[k] is None for r in grouped[f"{year}{m:02d}"])
                    else sum(r[k] for r in grouped[f"{year}{m:02d}"])
                )
                for k in ("a1", "a2", "c1")
            },
        }
        for m in range(1, 13)
    ]
