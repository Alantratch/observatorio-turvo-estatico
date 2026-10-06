import copy,csv,io,json,tempfile,unittest,zipfile
from pathlib import Path
from unittest.mock import patch
from scripts.sources import inep
from scripts.modules import education as e
FIX=Path(__file__).parent/'fixtures'/'education'
ROOT=Path(__file__).resolve().parents[1]
def fixture(name):return json.loads((FIX/name).read_text())
def schools():return fixture('schools-2025.json')
class Sheet:
 def __init__(self,table):self.title=table['sheet'];self.values=iter([table['header'],*table['rows']])
class Book:
 def __init__(self,tables):self.tables=tables
 def __iter__(self):return iter(Sheet(t) for t in self.tables)
 def close(self):pass
def metrics(key):
 f=fixture(key+'.json');source={**f['source'],'id':e.source_id(f['source'])}
 with patch.object(inep,'workbooks',return_value=[(f['tables'][0]['member'],Book(f['tables']))]):return e.indicators('fixture',source)
def school_zip(path,rows,omit=None):
 common={'NU_ANO_CENSO','CO_MUNICIPIO','SG_UF','CO_ENTIDADE','NO_ENTIDADE','TP_DEPENDENCIA','TP_LOCALIZACAO'}
 with zipfile.ZipFile(path,'w') as z:
  for kind,extras in [('Escola',{'TP_SITUACAO_FUNCIONAMENTO',*(k for r in rows for k in r if k.startswith('IN_'))}),('Matricula',set(e.ENROLMENTS.values())),('Docente',{'QT_DOC_BAS'}),('Turma',{'QT_TUR_BAS'})]:
   if kind==omit:continue
   keys=sorted(common|extras);buf=io.StringIO();w=csv.DictWriter(buf,keys,extrasaction='ignore',delimiter=';');w.writeheader();w.writerows(rows);z.writestr(f'dados/Tabela_{kind}_2025_V2.csv',buf.getvalue().encode('cp1252'))
class CensusTests(unittest.TestCase):
 def test_real_turvo_counts_and_stage_reconciliation(self):
  records,_,_=e.aggregate_schools(schools(),2025);r=next(r for r in records if r['code']=='4127965' and r['network']=='Total')
  self.assertEqual(r['enrolments']['total'],3230);self.assertEqual(r['activeSchools'],17);self.assertEqual(r['classes'],216);self.assertEqual(r['teachingPosts'],290)
  self.assertEqual(r['enrolments']['infant'],713);self.assertEqual(r['enrolments']['initial'],1188);self.assertEqual(r['enrolments']['final'],756)
 def test_special_education_does_not_increase_basic_total(self):
  a=e.aggregate_schools(schools(),2025)[0];r=next(r for r in a if r['code']=='4127965' and r['network']=='Total');self.assertEqual(r['enrolments']['special'],187)
  self.assertEqual(r['enrolments']['total'],r['enrolments']['infant']+r['enrolments']['fundamental']+r['enrolments']['secondary']+r['enrolments']['eja'])
 def test_networks_and_official_location(self):
  rs=e.aggregate_schools(schools(),2025)[0];mun=next(r for r in rs if r['code']=='4127965' and r['network']=='Municipal');self.assertEqual(mun['activeSchools'],9);self.assertEqual(mun['enrolments']['total'],1793)
  total=next(r for r in rs if r['code']=='4127965' and r['network']=='Total');self.assertEqual([r['schools'] for r in total['locations']],[8,9]);self.assertEqual([r['enrolments'] for r in total['locations']],[2195,1035])
 def test_duplicate_identical_school_not_counted_twice(self):self.assertEqual(e.aggregate_schools(schools()+[schools()[0]],2025),e.aggregate_schools(schools(),2025))
 def test_conflicting_duplicate_is_rejected(self):
  with self.assertRaises(ValueError):e.aggregate_schools(schools()+[{**schools()[0],'QT_MAT_BAS':'9999'}],2025)
 def test_bad_school_codes_and_state(self):
  for field,value in [('SG_UF','SC'),('NU_ANO_CENSO','2024'),('TP_DEPENDENCIA','9'),('TP_LOCALIZACAO','9')]:
   rs=schools();rs[0][field]=value
   with self.assertRaises(ValueError):e.aggregate_schools(rs,2025)
 def test_infrastructure_denominator_never_imputes_missing(self):
  rs=schools();rs[0]['IN_INTERNET']='';infra=e.aggregate_schools(rs,2025)[2];r=next(r for r in infra if r['code']=='4127965' and r['network']=='Total' and r['id']=='internet')
  self.assertEqual(r['missing'],1);self.assertEqual(r['considered'],16);self.assertAlmostEqual(r['percent'],r['withResource']/16*100,places=2)
 def test_invalid_binary_resource_rejected(self):
  rs=schools();rs[0]['IN_INTERNET']='2'
  with self.assertRaises(ValueError):e.aggregate_schools(rs,2025)
 def test_split_csv_join_and_homonym_excluded(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/'censo.zip';school_zip(p,schools()+[fixture('homonym-school.json')]);rs=list(inep.census_rows(p,2025));self.assertEqual(len(rs),len(schools()));self.assertTrue(all(r['SG_UF']=='PR' for r in rs));self.assertEqual(e.aggregate_schools(rs,2025),e.aggregate_schools(schools(),2025))
 def test_missing_split_table_fails_closed(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/'censo.zip';school_zip(p,schools(),omit='Matricula')
   with self.assertRaises(ValueError):list(inep.census_rows(p,2025))
 def test_declared_regular_school_missing_enrolment_is_not_zero(self):
  rs=schools();rs[0]['QT_MAT_BAS']='';rs[0]['IN_REGULAR']='1'
  with self.assertRaises(ValueError):e.aggregate_schools(rs,2025)
class IndicatorTests(unittest.TestCase):
 def test_suppressed_never_becomes_zero(self):
  for value in ('--','-',None,'ND','*',''):self.assertIsNone(e.numeric(value))
  self.assertEqual(e.numeric(0),0);self.assertEqual(e.numeric('7,4','IDEB'),7.4)
 def test_invalid_values_and_ideb_scale(self):
  for value,unit in [('nan','%'),('inf','%'),(-1,'%'),(101,'%'),(11,'IDEB')]:
   with self.assertRaises(ValueError):e.numeric(value,unit)
 def test_ideb_stage_network_cycle_and_homonym(self):
  rs=metrics('ideb');target=[r for r in rs if r['code']=='4127965' and r['network']=='Municipal' and r['reference']=='2025'];self.assertEqual(target[0]['value'],7.4);self.assertEqual(target[0]['stage'],'Anos iniciais');self.assertTrue(all(r['code']!='4218806' for r in rs));self.assertTrue(any(r['reference']=='2005' for r in rs));self.assertTrue(any(r['value'] is None for r in rs))
 def test_saeb_territorial_universes_do_not_collapse(self):
  self.assertEqual(e.network('Total - Estadual e Municipal'),'Estadual e Municipal');self.assertEqual(e.network('Total - Federal, Estadual e Municipal'),'Pública');self.assertEqual(e.network('Total - Federal, Estadual, Municipal e Privada'),'Total')
  rs=metrics('saeb');keys=[tuple(r[k] for k in ('code','network','location','stage','metric','reference')) for r in rs];self.assertEqual(len(keys),len(set(keys)))
 def test_literacy_has_official_target_and_participation(self):
  rs=[r for r in metrics('literacy') if r['code']=='4127965'];self.assertEqual(next(r['value'] for r in rs if r['metric']=='literacy' and r['reference']=='2025'),92);self.assertEqual(next(r['value'] for r in rs if r['metric']=='literacyParticipation'),96)
 def test_flow_same_universe_sums_to_100(self):
  rs=[r for r in metrics('flow') if r['code']=='4127965' and r['network']=='Total' and r['location']=='Total' and r['stage']=='Anos iniciais'];self.assertEqual(len(rs),3);self.assertAlmostEqual(sum(r['value'] for r in rs),100)
 def test_docente_groups_remain_separate(self):
  rs=[r for r in metrics('afd') if r['code']=='4127965' and r['network']=='Total' and r['location']=='Total' and r['stage']=='Anos iniciais'];self.assertEqual({r['metric'] for r in rs},{'afd1','afd2','afd3','afd4','afd5'});self.assertEqual(rs[0]['value'],86.7)
 def test_all_published_indicator_layouts_are_supported(self):
  for key in ('dsu','atu','had','tdi','ird','ied'):self.assertTrue(metrics(key),key)
 def test_official_uf_duplicate_unused_rendimento_header(self):
  s=Sheet({'sheet':'UF','header':['NO_UF','REDE','VL_OBSERVADO_2025','VL_INDICADOR_REND_2023','VL_INDICADOR_REND_2023'],'rows':[]})
  # The special header is accepted only with the actual human territorial header.
  s.values=iter([['Região/\nUnidade da Federação','Rede',None,None,None],[None,None,'VL_OBSERVADO_2025','VL_INDICADOR_REND_2023','VL_INDICADOR_REND_2023'],['Paraná','Pública',6.9,0,0]])
  self.assertEqual(list(e.machine_rows(s,True))[0]['VL_OBSERVADO_2025'],6.9)
class SnapshotTests(unittest.TestCase):
 def load(self):return json.loads((ROOT/'public/data/education.json').read_text())
 def test_snapshot_schema_all_sources_and_official_controls(self):
  d=self.load();e.validate(d);self.assertEqual(d['validation']['status'],'matched');self.assertEqual(len(d['validation']['officialChecks']),60);self.assertFalse(d['collection']['failures']);self.assertEqual(d['summary']['teachersOfficial'],250)
 def test_snapshot_has_eleven_census_years(self):
  years={r['reference'] for r in self.load()['census']['series'] if r['code']=='4127965' and r['network']=='Total'};self.assertEqual(years,{str(y) for y in range(2015,2026)})
 def test_failure_preserves_last_valid_numbers_and_sources(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t);d=self.load();(p/'education.json').write_text(json.dumps(d))
   with patch.object(inep,'discover',side_effect=OSError('rede indisponível')):out=e.update(p)
   self.assertEqual(out['census'],d['census']);self.assertEqual(out['indicators'],d['indicators']);self.assertTrue(out['collection']['failures'])
 def test_same_publication_never_downloads_again(self):
  d=self.load();m={}
  for s in d['sources']:m.setdefault(s['key'],[]).append(s)
  with tempfile.TemporaryDirectory() as t:
   p=Path(t);(p/'education.json').write_text(json.dumps(d))
   with patch.object(inep,'discover',return_value=m),patch.object(inep,'download',side_effect=AssertionError('download desnecessário')):out=e.update(p)
   self.assertFalse(out['collection']['failures']);self.assertEqual(out,d)
 def test_corrupt_infrastructure_and_status_rejected(self):
  for mutation in ('denominator','percent','status'):
   d=self.load()
   if mutation=='denominator':d['infrastructure'][0]['considered']+=1
   elif mutation=='percent':d['infrastructure'][0]['percent']=12.345
   else:d['indicators'][0]['status']='invented'
   with self.assertRaises(ValueError):e.validate(d)
 def test_corrupt_percentage_school_duplicate_or_schema_rejected(self):
  for mutation in ('schema','school','percent'):
   d=self.load()
   if mutation=='schema':d['schemaVersion']=2
   elif mutation=='school':d['schools'].append(d['schools'][0])
   else:next(r for r in d['indicators'] if r['unit']=='%')['value']=101
   with self.assertRaises(ValueError):e.validate(d)
 def test_unique_docentes_are_read_not_summed_across_schools(self):
  d=self.load();self.assertEqual(d['summary']['teachersOfficial'],250);self.assertEqual(d['summary']['teachingPosts'],290);self.assertNotEqual(d['summary']['teachersOfficial'],d['summary']['teachingPosts'])
class TransportTests(unittest.TestCase):
 def test_truncated_download_preserves_previous_file(self):
  class Response(io.BytesIO):
   headers={'Content-Length':'100','ETag':None,'Last-Modified':None}
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/'previous.zip';p.write_bytes(b'previous valid snapshot archive')
   with patch.object(inep,'request',return_value=Response(b'truncated')):
    with self.assertRaisesRegex(ValueError,'incompleto'):inep.download('https://download.inep.gov.br/test.zip',p,attempts=1)
   self.assertEqual(p.read_bytes(),b'previous valid snapshot archive');self.assertFalse(p.with_suffix('.zip.part').exists())
 def test_verified_cache_does_not_contact_remote(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/'cached.zip';p.write_bytes(b'verified cached archive');meta={'url':'https://download.inep.gov.br/test.zip','bytes':p.stat().st_size,'sha256':inep.file_hash(p)};p.with_suffix('.zip.manifest.json').write_text(json.dumps(meta))
   with patch.object(inep,'request',side_effect=AssertionError('unexpected remote call')):self.assertEqual(inep.download(meta['url'],p),meta)
 def test_real_projected_xlsx_can_be_read_from_zip(self):
  from openpyxl import Workbook
  f=fixture('ideb.json');book=Workbook();sheet=book.active;sheet.title=f['tables'][0]['sheet'];sheet.append(f['tables'][0]['header'])
  for row in f['tables'][0]['rows']:sheet.append(row)
  with tempfile.TemporaryDirectory() as t:
   xlsx=Path(t)/'projected.xlsx';book.save(xlsx);archive=Path(t)/'ideb.zip'
   with zipfile.ZipFile(archive,'w') as z:z.write(xlsx,f['tables'][0]['member'])
   rows=e.indicators(archive,{**f['source'],'id':e.source_id(f['source'])});self.assertEqual(rows,metrics('ideb'))
if __name__=='__main__':unittest.main()
